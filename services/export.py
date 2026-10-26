"""Builds a Word (.docx) file of one meeting: details, summary, decisions, action items, transcript.

Called by the History page's download button and process_cli.py --export.
"""
import io
import re
import unicodedata

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from services.store import get_meeting

NOTE = ("Owner is shown only when the conversation names them. Wording comes from an "
        "automatic transcript and may contain recognition errors.")
MISSING_ANALYSIS = ("The analysis is missing for this meeting (it failed). "
                    "Only the transcript is included.")
HEADER_SHADING = "D9E2F3"  # light blue behind the table header row

# XML 1.0 forbids most control characters (tab, newline and carriage return are allowed);
# python-docx raises if it gets one.
_NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f￾￿\ud800-\udfff]")


def clean_text(text) -> str:
    """Make any value safe to write into a .docx (removes characters XML does not allow)."""
    return _NOT_XML.sub("", "" if text is None else str(text))


def format_time(seconds) -> str:
    """mm:ss, or h:mm:ss from one hour up."""
    seconds = int(seconds or 0)
    hours, rest = divmod(seconds, 3600)
    minutes, secs = divmod(rest, 60)
    return f"{hours}:{minutes:02d}:{secs:02d}" if hours else f"{minutes:02d}:{secs:02d}"


def export_filename(title: str, created_at: str) -> str:
    """Safe file name: letters, digits, dash and underscore only, plus the date."""
    ascii_title = unicodedata.normalize("NFKD", title or "").encode("ascii", "ignore").decode()
    slug = re.sub(r"[^A-Za-z0-9_-]+", "-", ascii_title).strip("-_")
    slug = re.sub(r"-{2,}", "-", slug)[:60].rstrip("-_") or "meeting"
    return f"{slug}_{(created_at or '')[:10] or 'undated'}.docx"


def _set_up_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)  # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(25))
    doc.styles["Normal"].font.size = Pt(11)


def _shade(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _repeat_as_header(row) -> None:
    """Word repeats this row at the top of every page the table runs onto."""
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(marker)


def _action_table(doc: Document, items: list[dict]) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    header = table.rows[0]
    for cell, label in zip(header.cells, ("Task", "Owner", "Due", "Time")):
        cell.paragraphs[0].add_run(label).bold = True  # a new cell has one empty paragraph
        _shade(cell, HEADER_SHADING)
    _repeat_as_header(header)
    for item in items:
        cells = table.add_row().cells
        values = (item["task"], item["owner"] or "—", item["due_date"] or "—",
                  format_time(item["timestamp"]))
        for cell, value in zip(cells, values):
            cell.text = clean_text(value)


def _bullets(doc: Document, lines: list[str]) -> None:
    for line in lines or ["None."]:
        doc.add_paragraph(clean_text(line), style="List Bullet")


def build_docx(user_id: int, meeting_id: int) -> bytes | None:
    """The meeting as .docx bytes.

    None if the meeting is not found, belongs to someone else, or has no transcript.
    """
    meeting = get_meeting(user_id, meeting_id)
    if meeting is None or meeting["transcript"] is None:
        return None
    transcript, analysis = meeting["transcript"], meeting["analysis"]

    doc = Document()
    _set_up_page(doc)
    doc.add_heading(clean_text(meeting["title"]), level=1)
    created = meeting["created_at"][:16].replace("T", " ")
    for line in (f"Date: {created} UTC",
                 f"Language: {transcript.get('language') or 'unknown'}",
                 f"Duration: {format_time(meeting['duration_seconds'])}",
                 f"Source file: {meeting['source_filename']}"):
        doc.add_paragraph(clean_text(line))

    if analysis is None:
        doc.add_paragraph(MISSING_ANALYSIS).runs[0].bold = True
    else:
        summary, extraction = analysis["summary"], analysis["extraction"]
        doc.add_heading("Summary", level=2)
        doc.add_paragraph(clean_text(summary["overview"]))
        doc.add_paragraph("Key topics:")
        _bullets(doc, summary["key_topics"])

        doc.add_heading("Decisions", level=2)
        _bullets(doc, [f"[{format_time(d['timestamp'])}] {d['text']}"
                       for d in extraction["decisions"]])

        doc.add_heading("Action items", level=2)
        if extraction["action_items"]:
            _action_table(doc, extraction["action_items"])
        else:
            doc.add_paragraph("None.")

        doc.add_heading("Open topics", level=2)
        _bullets(doc, extraction["open_topics"])

        doc.add_paragraph().add_run(NOTE).italic = True

    doc.add_heading("Transcript", level=2)
    for seg in transcript["segments"]:
        doc.add_paragraph(clean_text(f"[{format_time(seg['start'])}] {seg['text'].strip()}"))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
