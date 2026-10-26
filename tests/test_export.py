"""Word export. Temp database, no API calls.
Run by pytest (uv run pytest).
"""
import io

import pytest
from docx import Document

from services import store
from services.export import (MISSING_ANALYSIS, NOTE, build_docx, clean_text, export_filename,
                             format_time)

SEGMENTS = [
    {"id": 0, "start": 0.0, "end": 5.0, "text": " They want $18,000, up from $15,000."},
    {"id": 1, "start": 64.0, "end": 70.0, "text": " Réunion à Pune — budget ₹2,50,000."},
    {"id": 2, "start": 3725.0, "end": 3730.0, "text": " Bad\x00 chars\x07 here\x1b."},
]
RESULT = {"model": "m",
          "summary": {"overview": "Vendor cap at $16,000; offsite in Pune.",
                      "key_topics": ["Vendor contract", "Offsite"]},
          "extraction": {
              "decisions": [{"topic": "price", "text": "Cap is $16,000.", "timestamp": 64}],
              "action_items": [
                  {"task": "Email the vendor", "owner": "Anita", "due_date": "November 8th",
                   "timestamp": 72},
                  {"task": "Book the venue\x0c", "owner": None, "due_date": None,
                   "timestamp": 3725}],
              "open_topics": ["Offsite dates"]},
          "steps": []}


@pytest.fixture
def users(temp_db):
    store.create_user("alice", "h", "user")
    store.create_user("bob", "h", "user")
    return store.get_user_by_username("alice")["id"], store.get_user_by_username("bob")["id"]


def meeting(user_id: int, status: str = "done") -> int:
    mid = store.create_meeting(user_id, "Vendor sync", "vendor.mp4")
    if status != "failed":
        store.save_transcript(user_id, mid, SEGMENTS, "whisper", 3730.0, language="English")
    if status == "done":
        store.save_analysis(user_id, mid, RESULT)
    else:
        store.update_status(user_id, mid, status, "something failed")
    return mid


def read(data: bytes) -> tuple[Document, str, list[list[str]]]:
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    rows = [[c.text for c in row.cells] for t in doc.tables for row in t.rows]
    return doc, text, rows


def test_done_meeting_has_everything(users):
    alice, _ = users
    doc, text, rows = read(build_docx(alice, meeting(alice)))
    assert doc.paragraphs[0].text == "Vendor sync"
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert "Language: English" in text and "Duration: 1:02:10" in text and " UTC" in text
    assert "Vendor cap at $16,000; offsite in Pune." in text
    assert "Vendor contract" in text and "Offsite" in text
    assert "[01:04] Cap is $16,000." in text
    assert rows == [["Task", "Owner", "Due", "Time"],
                    ["Email the vendor", "Anita", "November 8th", "01:12"],
                    ["Book the venue", "—", "—", "1:02:05"]]
    assert "Offsite dates" in text and NOTE in text
    assert "[00:00] They want $18,000, up from $15,000." in text
    assert "[01:04] Réunion à Pune — budget ₹2,50,000." in text
    assert "[1:02:05] Bad chars here." in text  # control characters removed, no crash
    section_headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert section_headings == ["Summary", "Decisions", "Action items", "Open topics", "Transcript"]


def test_table_and_page_format(users):
    alice, _ = users
    doc, _, _ = read(build_docx(alice, meeting(alice)))
    table = doc.tables[0]
    assert table.style.name == "Table Grid"
    header = table.rows[0]
    assert header._tr.trPr.xml.count("tblHeader") == 1  # repeated on every page
    assert all(c.paragraphs[0].runs[0].bold for c in header.cells)
    assert all("D9E2F3" in c._tc.tcPr.xml for c in header.cells)  # shaded
    assert doc.styles["Normal"].font.size.pt == 11
    section = doc.sections[0]
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (210, 297)


def test_analysis_failed_exports_transcript_with_note(users):
    alice, _ = users
    _, text, rows = read(build_docx(alice, meeting(alice, "analysis_failed")))
    assert MISSING_ANALYSIS in text
    assert "Summary" not in text and rows == []
    assert "[00:00] They want $18,000, up from $15,000." in text


def test_no_transcript_or_other_user_gives_none(users):
    alice, bob = users
    assert build_docx(alice, meeting(alice, "failed")) is None
    assert build_docx(bob, meeting(alice)) is None
    assert build_docx(alice, 99999) is None


def test_clean_text():
    assert clean_text("a\x00b\x07c\x1bd\x7fe\tf\ng") == "abcde\tf\ng"
    assert clean_text(None) == ""
    assert clean_text("₹ é $") == "₹ é $"


def test_format_time():
    assert [format_time(s) for s in (0, 65, 3599, 3600, 3725)] == \
        ["00:00", "01:05", "59:59", "1:00:00", "1:02:05"]


@pytest.mark.parametrize("title, expected", [
    ("Q3 / Budget: $18,000", "Q3-Budget-18-000"),
    ("Réunion à Pune", "Reunion-a-Pune"),
    ("../../etc/passwd", "etc-passwd"),
    ("weekly_sync-2", "weekly_sync-2"),
    ("", "meeting"),
    ("🎉🎉", "meeting"),
    ("x" * 100, "x" * 60),
])
def test_export_filename(title, expected):
    assert export_filename(title, "2026-10-25T15:12:00+00:00") == f"{expected}_2026-10-25.docx"
